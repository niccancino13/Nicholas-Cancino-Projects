import ipaddress
import re
from openpyxl import load_workbook

# Pool members are exported as "server-name(10.x.x.x)" strings
POOL_MEMBER_RE = re.compile(r"([A-Za-z0-9._-]+)\((\d{1,3}(?:\.\d{1,3}){3})\)")
VIP_COL, POOL_COL = 2, 4  # zero-based column indexes in the F5 export

def is_external(ip_str: str) -> bool:
    """True only for publicly routable addresses (drops RFC 1918, loopback, reserved)."""
    try:
        return ipaddress.ip_address(ip_str).is_global
    except ValueError:
        return False

def external_vips(report_path: str):
    ws = load_workbook(report_path, read_only=True).active
    for row in ws.iter_rows(min_row=2, values_only=True):
        # VIP column, e.g. "x.x.x.x:443" or "x.x.x.x%1:443" (adjust index to your export)
        vip = str(row[VIP_COL] or "").split("%")[0].split(":")[0]
        if not is_external(vip):
            continue
        members = POOL_MEMBER_RE.findall(str(row[POOL_COL] or ""))  # pool member column
        yield vip, members  # members -> [(server_name, server_ip), ...]
